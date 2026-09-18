"""卡片配置的表格编解码层：CSV / XLSX ⇄ list[dict]。

为什么单独抽一层
----------------
JSON 导入导出的"记录形状"由 `CardExportItem` 定义。表格只是同一份记录的另一
种载体，所以这一层**只做编解码、不做业务校验**：

    bytes ──(parse_sheet)──> list[dict] ──(CardService.import_cards)──> 落库

校验、覆盖/合并语义、错误收集全部复用 `CardService`，因此"表格能填什么"和
"接口接受什么"永远是同一份规则，不会出现表格允许、接口拒绝的错位。

几个刻意的取舍
--------------
1. **表头写中文，解析同时认英文**：导出的文件是给人看的，中文表头才友好；
   但手工攒的文件常写 `slug`/`target_url`，所以别名表两套都收。匹配时忽略
   大小写、前后空白、全角空格、下划线与连字符。
2. **CSV 导出带 UTF-8 BOM**：不带 BOM 时 Excel(Windows) 会按 ANSI/GBK 解码，
   中文描述直接变乱码。导入时依次尝试 UTF-16 / UTF-8(BOM) / GB18030。
3. **空单元格 = 用默认值**：与本项目 `CardExportItem` 的默认值语义一致。
   注意这意味着"整行覆盖"——合并导入时，留空的列会被重置为默认值，
   所以 UI 上必须先跑一次预演把影响摆出来。
4. **`verify_tls` 是三态**：留空 = None（跟随全局），不是 False。布尔解析里
   空值必须单独处理，否则会把"不校验"错误地写成"校验"。
5. **非法单元格只废掉那一行**：抛 `_CellError` 而不是 `BadRequestError`，
   由调用方记入错误清单并跳过该行，与 JSON 导入的行为保持一致。
"""

from __future__ import annotations

import csv
import io
import json
import logging
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.errors import BadRequestError
from app.models.entities import OPEN_MODE_DIRECT
from app.schemas.models import BG_STYLES, CardExportDoc

logger = logging.getLogger("app.service.card_sheet")

# 单次导入的数据行上限，与 CardImportRequest.cards 的 max_length 保持一致
MAX_DATA_ROWS = 5000
# 行级错误最多保留多少条（最终展示还会被 CardService 再截一次）
MAX_ERRORS = 100
# 表头最多在前几行内找（兼容用户自己在上面加了一行标题）
HEADER_SCAN_ROWS = 10
# 错误提示里单元格内容最多展示多少字符
CELL_PREVIEW = 40

_XLSX_SHEET_CARDS = "卡片"
_XLSX_SHEET_HELP = "字段说明"

# 文字图标配色：与前端 frontend/src/lib/appearance.ts 的 ICON_STYLES 保持一致
ICON_STYLE_OPTIONS: tuple[str, ...] = (
    "blue",
    "indigo",
    "violet",
    "cyan",
    "teal",
    "green",
    "amber",
    "rose",
)

OPEN_MODE_OPTIONS: tuple[str, ...] = ("direct", "proxy")

# Excel 的下拉列表里不能有"空选项"，所以把 BG_STYLES 里的空串去掉
BG_STYLE_OPTIONS: tuple[str, ...] = tuple(style for style in BG_STYLES if style)

BOOL_HINT = "可填 是/否、true/false、1/0、y/n"

# 「环境地址」列在表格里的写法：一格装多条，每条 `名称 | 地址 | 模式`。
# 之所以塞进一格而不是给 xlsx 单开一张工作表：CSV 只有一张表，一旦为 xlsx 设计
# 第二张表，CSV 就表达不了环境地址，"三种格式都能导出导入"的承诺就破了。
_ENDPOINTS_KEY = "endpoints"
_ENDPOINT_DELIMITER = "|"
_ENDPOINT_MODE_ALIASES: dict[str, str] = {
    # _norm_header 会去掉空白/下划线/连字符并转小写，所以这里写归一化之后的键
    "direct": "direct",
    "直连": "direct",
    "直接": "direct",
    "直连访问": "direct",
    "跳转": "direct",
    "proxy": "proxy",
    "代理": "proxy",
    "反向代理": "proxy",
    "跳板机": "proxy",
}


class _CellError(ValueError):
    """单元格取值无法识别。

    只跳过该行、不影响整批导入，因此不用 `BadRequestError`（那会整单 400）。
    """


@dataclass(frozen=True)
class SheetColumn:
    """表格的一列：列定义是导出与导入共用的唯一事实来源。"""

    key: str
    header: str
    note: str
    example: str = ""
    required: bool = False
    options: tuple[str, ...] = ()
    # text | bool | bool_or_blank | int
    kind: str = "text"
    width: int = 14
    aliases: tuple[str, ...] = ()
    # 取值别名：表头是中文，用户很自然会顺手把取值也写成中文（直连 / 天青），
    # 这里把它们映射回官方取值，避免"看得懂表头却填不对值"
    value_aliases: tuple[tuple[str, str], ...] = ()


CARD_COLUMNS: tuple[SheetColumn, ...] = (
    SheetColumn(
        key="slug",
        header="标识",
        note="卡片的唯一标识。导入时按它匹配已有卡片：命中就更新，没命中就新建。只能用小写字母、数字和连字符。",
        example="jenkins",
        required=True,
        width=16,
        aliases=("id", "key", "identifier", "唯一标识"),
    ),
    SheetColumn(
        key="title",
        header="名称",
        note="卡片上显示的名字。",
        example="Jenkins",
        required=True,
        width=18,
        aliases=("标题", "name", "cardtitle"),
    ),
    SheetColumn(
        key="group_name",
        header="分组",
        note="首页按分组归类展示。留空按「默认分组」处理。",
        example="CI/CD",
        width=14,
        aliases=("分组名", "group", "category", "所属分组"),
    ),
    SheetColumn(
        key="target_url",
        header="目标地址",
        note="卡片要打开的内网地址，必须以 http:// 或 https:// 开头。以 / 结尾表示目录作用域。",
        example="http://10.0.0.20:8081/",
        required=True,
        width=46,
        aliases=("地址", "链接", "url", "target", "targeturl", "跳转地址"),
    ),
    SheetColumn(
        key="open_mode",
        header="打开方式",
        note="direct = 浏览器直连目标；proxy = 经跳板机反向代理（同源，登录态会与门户互相影响）。",
        example="direct",
        options=OPEN_MODE_OPTIONS,
        width=12,
        aliases=("模式", "mode", "openmode", "访问方式"),
        value_aliases=(
            ("直连", "direct"),
            ("直接", "direct"),
            ("直连访问", "direct"),
            ("代理", "proxy"),
            ("反向代理", "proxy"),
            ("跳板机", "proxy"),
        ),
    ),
    SheetColumn(
        key="endpoints",
        header="环境地址",
        note=(
            "同一张卡片可以挂多条地址（生产 / 测试 / 预发 / 灰度…），点卡片时会弹出列表让人选。\n"
            "每条写成「名称 | 地址 | 模式」，多条之间用换行分隔（Excel 里按 Alt+Enter 换行），"
            "也可以直接用分号隔开。模式可省略，省略时跟随本行的「打开方式」；本行也没填就按 direct。\n"
            "这一列留空 = 这张卡片没有环境地址（导入时会清空已配置的）；"
            "文件里根本没有这一列时，库里的环境地址会原样保留。"
        ),
        example="生产 | https://jenkins.example.com/\n测试 | http://10.0.0.9:8081/ | proxy",
        width=56,
        aliases=("环境", "多地址", "地址列表", "环境列表", "endpoints", "endpoint", "envs"),
        kind="endpoints",
    ),
    SheetColumn(
        key="description",
        header="描述",
        note="鼠标悬停在卡片上时显示的说明文字。",
        example="持续集成流水线",
        width=40,
        aliases=("说明", "描述信息", "desc", "remark", "备注"),
    ),
    SheetColumn(
        key="icon",
        header="图标文字",
        note="没有上传图标图片时，用 1~2 个字符当图标。留空则取名称首字。",
        example="J",
        width=11,
        aliases=("图标字", "icontext", "图标"),
    ),
    SheetColumn(
        key="icon_style",
        header="图标配色",
        note="文字图标的配色预设。",
        example="blue",
        options=ICON_STYLE_OPTIONS,
        width=12,
        aliases=("图标颜色", "iconstyle", "配色"),
        value_aliases=(
            ("天蓝", "blue"),
            ("靛蓝", "indigo"),
            ("紫罗兰", "violet"),
            ("青色", "cyan"),
            ("蓝绿", "teal"),
            ("绿色", "green"),
            ("琥珀", "amber"),
            ("玫红", "rose"),
        ),
    ),
    SheetColumn(
        key="icon_url",
        header="图标图片",
        note="填 /uploads/... 本站图片或 http(s) 外链；留空表示用上面的图标文字。",
        example="/uploads/icons/8f3c....png",
        width=30,
        aliases=("图标地址", "iconurl", "图片"),
    ),
    SheetColumn(
        key="bg_style",
        header="卡片底色",
        note=f"卡片背景配色预设，留空为默认蓝色渐变。可选：{'、'.join(BG_STYLE_OPTIONS)}。",
        example="blue",
        options=BG_STYLE_OPTIONS,
        width=12,
        aliases=("底色", "背景", "bgstyle", "cardbg"),
        value_aliases=(
            ("默认", ""),
            ("云雾蓝", "default"),
            ("天青", "blue"),
            ("紫罗兰", "violet"),
            ("薄荷", "mint"),
            ("暖沙", "sand"),
            ("浅玫", "rose"),
            ("石板", "slate"),
        ),
    ),
    SheetColumn(
        key="accent_color",
        header="图标主色",
        note="自定义图标主色，必须写成 #RRGGBB。留空表示跟随卡片底色。",
        example="#2b6cf6",
        width=12,
        aliases=("主色", "accent", "accentcolor", "颜色"),
    ),
    SheetColumn(
        key="enabled",
        header="启用",
        note="填「否」则该卡片不在首页显示（管理员在后台仍能看到）。留空按「是」处理。",
        example="是",
        kind="bool",
        width=9,
        aliases=("是否启用", "启用状态", "on", "active"),
    ),
    SheetColumn(
        key="open_in_new_tab",
        header="新标签页打开",
        note="「是」= 新标签页；「否」= 在当前标签页跳转。留空按「是」处理。",
        example="是",
        kind="bool",
        width=15,
        aliases=("新窗口打开", "openinnewtab", "newtab", "新标签页"),
    ),
    SheetColumn(
        key="verify_tls",
        header="校验HTTPS证书",
        note="留空 = 跟随服务端全局配置；「否」= 跳过证书校验（内网自签证书常见）。仅代理模式生效。",
        example="",
        kind="bool_or_blank",
        width=15,
        aliases=("校验证书", "verifyssl", "verifytls", "跳过证书校验"),
    ),
    SheetColumn(
        key="sort_order",
        header="排序",
        note="数字越小越靠前。留空或填 0 表示按导入顺序自动排到末尾。",
        example="10",
        kind="int",
        width=9,
        aliases=("顺序", "排序值", "sort", "order", "sortorder", "排序号"),
    ),
)

_COLUMN_BY_KEY: dict[str, SheetColumn] = {column.key: column for column in CARD_COLUMNS}
_REQUIRED_KEYS: tuple[str, ...] = tuple(
    column.key for column in CARD_COLUMNS if column.required
)

# 布尔取值：中英文日常写法都收，避免"填了是却说格式不对"
_TRUE_WORDS = {
    "1", "y", "yes", "true", "t", "on", "enable", "enabled",
    "是", "对", "有", "真", "开", "启用", "开启", "勾选", "✓", "✔", "√",
}
_FALSE_WORDS = {
    "0", "n", "no", "false", "f", "off", "disable", "disabled",
    "否", "不", "无", "假", "关", "停用", "关闭", "取消", "x", "×", "✗",
}


@dataclass
class SheetParseResult:
    """表格解析结果。

    `cards` / `row_numbers` 一一对应：`row_numbers[i]` 是 `cards[i]` 在表格里的
    物理行号（CSV 行号 / Excel 行号），用于把错误提示定位到用户看得见的位置。
    """

    cards: list[Any]
    row_numbers: list[int]
    errors: list[str]
    total: int
    header_row: int = 0


# ----------------------------------------------------------------------
# 列名匹配
# ----------------------------------------------------------------------
def _norm_header(value: object) -> str:
    """把表头归一化：去掉空白/全角空格/下划线/连字符并转小写。"""
    text = "" if value is None else str(value)
    for char in ("\u3000", " ", "\t", "\n", "\r", "_", "-", "/", "／", "（", "）"):
        text = text.replace(char, "")
    return text.strip().lower()


def _build_alias_map() -> dict[str, str]:
    out: dict[str, str] = {}
    for column in CARD_COLUMNS:
        for alias in (column.header, column.key, *column.aliases):
            normalized = _norm_header(alias)
            if normalized:
                out.setdefault(normalized, column.key)
    return out


_ALIAS_MAP: dict[str, str] = _build_alias_map()


def _resolve_header(
    rows: Sequence[tuple[int, list[Any]]],
) -> tuple[int, dict[int, str]]:
    """在前若干行里找表头，返回 (表头在 rows 里的下标, {列下标: 字段名})。"""
    best_index = -1
    best_map: dict[int, str] = {}
    for index, (_row_no, cells) in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        taken: set[str] = set()
        for position, cell in enumerate(cells):
            key = _ALIAS_MAP.get(_norm_header(cell))
            # 同名列只认最左边那一列，避免后一列把前一列覆盖掉
            if key and key not in taken:
                taken.add(key)
                mapping[position] = key
        if len(mapping) > len(best_map):
            best_index, best_map = index, mapping
    if len(best_map) < 2:
        return -1, {}
    return best_index, best_map


# ----------------------------------------------------------------------
# 取值转换
# ----------------------------------------------------------------------
def _cell_text(value: Any) -> str:
    """把单元格原始值转成字符串（兼容 Excel 自动识别的日期/数字/布尔）。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return str(value).strip()


def _preview(text: str) -> str:
    flat = " ".join(text.split())
    return flat if len(flat) <= CELL_PREVIEW else flat[:CELL_PREVIEW] + "…"


def _parse_bool(text: str) -> bool:
    key = _norm_header(text)
    if key in _TRUE_WORDS:
        return True
    if key in _FALSE_WORDS:
        return False
    raise _CellError(f"的值「{_preview(text)}」无法识别（{BOOL_HINT}）")


def _match_option(text: str, column: SheetColumn) -> str | None:
    """把单元格文本匹配到列定义的官方取值；认不出来返回 None。"""
    key = _norm_header(text)
    for alias, canonical in column.value_aliases:
        if _norm_header(alias) == key:
            return canonical
    for option in column.options:
        if _norm_header(option) == key:
            return option
    return None


def _convert(value: Any, column: SheetColumn) -> Any:
    """把单元格值转成字段值；返回 None 表示"这格是空的"，调用方会省略该字段。"""
    text = _cell_text(value)
    if not text:
        return None
    if column.options:
        # 枚举列在这里就判掉，否则错误会落到 Pydantic 手里变成英文的
        # "Input should be 'direct' or 'proxy'"，用户看不懂该填什么
        matched = _match_option(text, column)
        if matched is None:
            raise _CellError(
                f"的值「{_preview(text)}」不在允许范围内（可填：{'、'.join(column.options)}）"
            )
        return matched
    if column.kind in ("bool", "bool_or_blank"):
        return _parse_bool(text)
    if column.kind == "int":
        try:
            return int(float(text))
        except (ValueError, OverflowError) as exc:
            raise _CellError(f"的值「{_preview(text)}」不是数字") from exc
    return text


def _split_endpoint_entries(text: str) -> list[str]:
    """把一格里的多条环境地址拆开。

    优先按换行拆；**整格没有换行时**才退回按分号拆。这样安排的理由是
    地址里合法地可能出现分号（`/path;a=b` 这种矩阵参数），只要用户是
    "一条一行"（Excel 里 Alt+Enter 很正常）就绝不会被拆坏。
    """
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    parts = normalized.split("\n") if "\n" in normalized else normalized.split(";")
    return [part.strip() for part in parts if part.strip()]


def _parse_endpoints(text: str, default_mode: str) -> list[dict[str, Any]]:
    """把「环境地址」单元格解析成 `CardEndpointIn` 能吃的字典列表。

    认不出来就抛 `_CellError`，由调用方废掉那一行并给出带行号的提示 ——
    绝不"猜一半、丢一半"：少一条归档地址比报错难发现得多。
    """
    items: list[dict[str, Any]] = []
    for index, entry in enumerate(_split_endpoint_entries(text), start=1):
        # 全角竖线也认，中文输入法下打出来的是这个
        fields = [field.strip() for field in entry.replace("｜", _ENDPOINT_DELIMITER).split(_ENDPOINT_DELIMITER)]
        if len(fields) > 3:
            raise _CellError(
                f"的第 {index} 条「{_preview(entry)}」多了分隔符，"
                "每条只能是「名称 | 地址 | 模式」三段（模式可省略）"
            )
        if len(fields) < 2 or not fields[0] or not fields[1]:
            raise _CellError(
                f"的第 {index} 条「{_preview(entry)}」格式不对，应写成「名称 | 地址 | 模式」（模式可省略）"
            )
        name, url = fields[0], fields[1]
        if not url.lower().startswith(("http://", "https://")):
            raise _CellError(
                f"的第 {index} 条地址「{_preview(url)}」必须以 http:// 或 https:// 开头"
            )
        mode = default_mode
        if len(fields) == 3 and fields[2]:
            matched = _ENDPOINT_MODE_ALIASES.get(_norm_header(fields[2]))
            if matched is None:
                raise _CellError(
                    f"的第 {index} 条模式「{_preview(fields[2])}」不在允许范围内（可填：direct、proxy）"
                )
            mode = matched
        items.append({"name": name, "url": url, "open_mode": mode})
    return items


def _format_endpoints(items: Any, separator: str) -> str:
    """把一条卡片记录里的环境地址格式化成单元格文本。"""
    parts: list[str] = []
    for entry in items or []:
        name = str(_field_value(entry, "name") or "").strip()
        url = str(_field_value(entry, "url") or "").strip()
        mode = str(_field_value(entry, "open_mode") or "").strip()
        if not name and not url:
            continue
        parts.append(f"{name} {_ENDPOINT_DELIMITER} {url} {_ENDPOINT_DELIMITER} {mode}")
    return separator.join(parts)


def _cell_value(item: Any, column: SheetColumn, *, multi_sep: str = "\n") -> str | int | None:
    """把一条卡片记录格式化成单元格值。None 表示写空单元格。

    `multi_sep` 是一格多条（目前只有环境地址列）时的分隔符：
    xlsx 用换行（配 wrap_text，一格能看全），CSV 用分号（保持"一行一条记录"，
    避免嵌入换行把一些第三方表格工具读崩）。
    """
    raw = item.get(column.key) if isinstance(item, Mapping) else getattr(item, column.key, None)

    if column.kind == "bool":
        return "是" if bool(raw) else "否"
    if column.kind == "bool_or_blank":
        if raw is None:
            return None  # 三态：空 = 跟随全局，不能写成「否」
        return "是" if bool(raw) else "否"
    if column.kind == "int":
        try:
            return int(raw or 0)
        except (TypeError, ValueError):
            return 0
    if column.kind == "endpoints":
        text = _format_endpoints(raw, multi_sep)
        return text or None
    text = "" if raw is None else str(raw)
    return text or None


# ----------------------------------------------------------------------
# 行 → 卡片字典
# ----------------------------------------------------------------------
_NO_COLUMN_HINT = (
    "没在文件里找到卡片列。表头需要至少包含「标识」「名称」「目标地址」三列，"
    "建议先用本工具导出一份表格当模板。"
)


def _rows_to_cards(
    rows: Sequence[tuple[int, list[Any]]], *, source: str
) -> SheetParseResult:
    header_index, mapping = _resolve_header(rows)
    if header_index < 0:
        raise BadRequestError(f"{source}：{_NO_COLUMN_HINT}")

    found = set(mapping.values())
    missing = [
        _COLUMN_BY_KEY[key].header for key in _REQUIRED_KEYS if key not in found
    ]
    if missing:
        raise BadRequestError(f"{source} 缺少必需的列：{'、'.join(missing)}")

    header_row = rows[header_index][0]
    cards: list[Any] = []
    row_numbers: list[int] = []
    errors: list[str] = []
    total = 0

    for row_no, cells in rows[header_index + 1 :]:
        raw = {key: cells[pos] for pos, key in mapping.items() if pos < len(cells)}
        # 整行空白（Excel 尾部常有）直接跳过，不算一行数据
        if all(_cell_text(value) == "" for value in raw.values()):
            continue

        total += 1
        if total > MAX_DATA_ROWS:
            raise BadRequestError(
                f"{source} 的数据行超过单次导入上限 {MAX_DATA_ROWS} 行，请拆成多个文件分批导入。"
            )

        card: dict[str, Any] = {}
        failed = False
        for key, value in raw.items():
            if key == _ENDPOINTS_KEY:
                # 单独处理：省略模式时默认值要参考同一行的「打开方式」，
                # 而字典的遍历顺序不保证 open_mode 已经先解析完
                continue
            column = _COLUMN_BY_KEY[key]
            try:
                converted = _convert(value, column)
            except _CellError as exc:
                if len(errors) < MAX_ERRORS:
                    errors.append(f"第 {row_no} 行：{column.header}{exc}，本行已跳过")
                failed = True
                break
            if converted is not None:
                card[key] = converted

        # 环境地址列「出现即权威」：
        #   有这一列 + 单元格留空  → 写成空列表，导入时清空该卡片已配置的地址
        #   文件里没这一列        → 键根本不进 card，导入时库里的地址原样保留
        # 这条区分是刻意留的：本功能上线前导出的老表格没有这一列，
        # 拿它回来"改个错别字"不该把用户攒下的归档地址一起抹掉。
        if not failed and _ENDPOINTS_KEY in raw:
            text = _cell_text(raw[_ENDPOINTS_KEY])
            try:
                card[_ENDPOINTS_KEY] = (
                    _parse_endpoints(text, card.get("open_mode") or OPEN_MODE_DIRECT)
                    if text
                    else []
                )
            except _CellError as exc:
                if len(errors) < MAX_ERRORS:
                    errors.append(
                        f"第 {row_no} 行：{_COLUMN_BY_KEY[_ENDPOINTS_KEY].header}{exc}，本行已跳过"
                    )
                failed = True

        if failed:
            continue
        cards.append(card)
        row_numbers.append(row_no)

    logger.info(
        "解析%s完成：表头第 %d 行，数据 %d 行，可用 %d 行，行级错误 %d 条",
        source, header_row, total, len(cards), len(errors),
    )
    return SheetParseResult(
        cards=cards,
        row_numbers=row_numbers,
        errors=errors,
        total=total,
        header_row=header_row,
    )


# ----------------------------------------------------------------------
# 文本解码
# ----------------------------------------------------------------------
def _decode_text(data: bytes) -> str:
    """按 BOM 优先、其次 UTF-8、最后 GB18030 的顺序解码文本。"""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BadRequestError(
        "无法识别文件的文本编码。请用「UTF-8」或「GBK/GB18030」重新保存后再导入。"
    )


def _sniff_delimiter(text: str) -> str:
    """判断 CSV 分隔符。

    只在我们自己的列语义下判断：取第一行非空文本，谁能切出 ≥3 列就用谁。
    不引入 `csv.Sniffer` —— 它在小样本上经常误判，而且报错方式不友好。
    """
    for line in text.splitlines():
        if not line.strip():
            continue
        for candidate in (",", ";", "\t", "|"):
            try:
                fields = next(csv.reader([line], delimiter=candidate))
            except csv.Error:
                continue
            if len(fields) >= 3:
                return candidate
        return ","
    return ","


# ----------------------------------------------------------------------
# 读：CSV / XLSX / JSON
# ----------------------------------------------------------------------
_ROW_READ_LIMIT = MAX_DATA_ROWS + HEADER_SCAN_ROWS + 2


def from_csv(data: bytes) -> SheetParseResult:
    text = _decode_text(data)
    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)

    rows: list[tuple[int, list[Any]]] = []
    try:
        for record in reader:
            rows.append((reader.line_num, list(record)))
            if len(rows) >= _ROW_READ_LIMIT:
                break
    except csv.Error as exc:
        raise BadRequestError(f"CSV 格式有误，无法解析：{exc}") from exc

    return _rows_to_cards(rows, source="CSV")


def _read_sheet_rows(sheet: Any, *, max_row: int) -> list[tuple[int, list[Any]]]:
    """读取工作表，返回 (Excel 行号, 该行各列的值)。"""
    rows: list[tuple[int, list[Any]]] = []
    for row in sheet.iter_rows(min_row=1, max_row=max_row):
        if not row:
            continue
        # 非只读模式下 Cell.row 是精确的物理行号，错误提示才能对上 Excel
        rows.append((row[0].row, [cell.value for cell in row]))
    return rows


def _pick_sheet(workbook: Any) -> tuple[Any, int]:
    """挑出卡片表：第一个"表头能认出 ≥3 列"的工作表优先。"""
    fallback: tuple[Any, int] | None = None
    for name in workbook.sheetnames:
        sheet = workbook[name]
        rows = _read_sheet_rows(sheet, max_row=HEADER_SCAN_ROWS)
        header_index, mapping = _resolve_header(rows)
        if len(mapping) >= 3:
            return sheet, rows[header_index][0]
        if fallback is None:
            fallback = (sheet, 1)
    # 都不像卡片表，用第一张表交给后续报错（错误信息会指明缺哪几列）
    assert fallback is not None
    return fallback


def from_xlsx(data: bytes) -> SheetParseResult:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except InvalidFileException as exc:
        raise BadRequestError(
            "无法识别这个 Excel 文件。旧版 .xls 请先用 Excel 另存为 .xlsx 或 .csv 再导入。"
        ) from exc
    except Exception as exc:  # noqa: BLE001 —— 解析不可信文件，边界处统一兜住
        logger.warning("读取 xlsx 失败: %s", exc)
        raise BadRequestError("Excel 文件已损坏，或不是有效的 .xlsx 文件。") from exc

    try:
        sheet, header_row = _pick_sheet(workbook)
        # 多读几行是为了在超限时能明确报错，而不是静默截断
        rows = _read_sheet_rows(sheet, max_row=_ROW_READ_LIMIT)
        result = _rows_to_cards(rows, source=f"工作表「{sheet.title}」")
        result.header_row = header_row
        return result
    finally:
        workbook.close()


def from_json(data: bytes) -> SheetParseResult:
    """JSON 不算"表格"，但上传接口统一收它，所以在这里做同构包装。"""
    try:
        raw = json.loads(_decode_text(data))
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise BadRequestError("文件不是合法的 JSON，请选择本工具导出的配置文件。") from exc

    if isinstance(raw, list):
        cards: Any = raw
    elif isinstance(raw, Mapping):
        cards = raw.get("cards")
    else:
        cards = None

    if not isinstance(cards, list):
        raise BadRequestError("JSON 里没有找到 cards 数组，请确认是本工具导出的配置文件。")

    # 不填 row_numbers：交给 CardService 报错时用「第 N 条」编号
    return SheetParseResult(cards=list(cards), row_numbers=[], errors=[], total=len(cards))


# ----------------------------------------------------------------------
# 格式识别与统一入口
# ----------------------------------------------------------------------
def _extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    return "." + name.rsplit(".", 1)[-1] if "." in name else ""


def detect_format(data: bytes, filename: str = "") -> str:
    """识别文件格式：返回 xlsx / xls / json / csv。

    以**内容**为主、扩展名为辅 —— 用户把 xlsx 改名成 .csv 是常见误操作，
    只看扩展名会给出莫名其妙的报错。
    """
    head = data[:8]
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if data[:3] == b"\xef\xbb\xbf":
        # 带 BOM 的 UTF-8：可能是 CSV，也可能是 Windows 上另存的 JSON
        return "json" if data[3:].lstrip()[:1] in (b"{", b"[") else "csv"
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "csv"

    stripped = data.lstrip()[:1]
    if stripped in (b"{", b"["):
        return "json"

    extension = _extension(filename)
    if extension == ".json":
        return "json"
    if extension in (".xlsx", ".xlsm"):
        return "xlsx"
    if extension == ".xls":
        return "xls"
    return "csv"


def parse_sheet(data: bytes, filename: str = "") -> SheetParseResult:
    """把任意受支持的导入文件解析成卡片字典列表。"""
    if not data:
        raise BadRequestError("文件是空的，没有可导入的内容。")

    kind = detect_format(data, filename)
    if kind == "xlsx":
        return from_xlsx(data)
    if kind == "xls":
        raise BadRequestError(
            "不支持旧版 .xls 格式。请用 Excel 打开后「另存为」.xlsx 或 .csv 再导入。"
        )
    if kind == "json":
        return from_json(data)
    # UTF-16 的文本里本来就到处是 0x00，只有非 UTF-16 才按二进制处理
    if data[:2] not in (b"\xff\xfe", b"\xfe\xff") and b"\x00" in data[:4096]:
        raise BadRequestError(
            "这个文件是二进制内容，不是卡片表格。请选择导出的 .xlsx / .csv / .json 文件。"
        )
    return from_csv(data)


# ----------------------------------------------------------------------
# 写：CSV / XLSX / JSON
# ----------------------------------------------------------------------
def _field_value(item: Any, key: str) -> Any:
    return item.get(key) if isinstance(item, Mapping) else getattr(item, key, None)


def to_json(doc: CardExportDoc) -> bytes:
    payload = doc.model_dump(mode="json")
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def to_csv(doc: CardExportDoc) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([column.header for column in CARD_COLUMNS])
    for item in doc.cards:
        writer.writerow(
            [
                "" if (value := _cell_value(item, column, multi_sep=";")) is None else value
                for column in CARD_COLUMNS
            ]
        )
    # 必须带 BOM：否则 Excel(Windows) 会按 GBK 解码，中文全变乱码
    return buffer.getvalue().encode("utf-8-sig")


def _style_cards_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="E8EFFB")
    header_font = Font(bold=True, color="1B2430")
    last_column = get_column_letter(len(CARD_COLUMNS))
    last_row = max(sheet.max_row, 1)

    for index, column in enumerate(CARD_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="left")
        note = column.note
        if column.example:
            note += f"\n示例：{column.example}"
        cell.comment = Comment(note, "开发工具箱")
        sheet.column_dimensions[get_column_letter(index)].width = column.width
        if column.kind == "endpoints":
            # 一格里塞多条地址，必须自动换行，否则用户只看得见第一条
            sheet.column_dimensions[get_column_letter(index)].alignment = Alignment(
                wrap_text=True, vertical="top"
            )

    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"  # 滚到第 500 行也能看见表头
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"

    # 枚举列加下拉框；多留 200 行，方便用户继续往下加卡片
    validation_range_end = last_row + 200
    for index, column in enumerate(CARD_COLUMNS, start=1):
        if not column.options:
            continue
        letter = get_column_letter(index)
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(column.options) + '"',
            allow_blank=True,
            # OOXML 里 showDropDown=True 的实际含义是"隐藏下拉箭头"，与名字相反
            showDropDown=False,
        )
        validation.errorTitle = "取值不在允许范围"
        validation.error = f"{column.header} 只能填：{'、'.join(column.options)}"
        validation.promptTitle = column.header
        validation.prompt = "可选值：" + " / ".join(column.options)
        validation.showErrorMessage = True
        validation.showInputMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{validation_range_end}")


_HELP_COLUMNS = ("字段", "含义", "是否必填", "取值示例")
# 注意别把说明页的表头写成卡片列的别名，否则挑工作表时可能误判
_HELP_WIDTHS = (16, 78, 10, 26)


def _write_help_sheet(workbook: Workbook, doc: CardExportDoc) -> None:
    sheet = workbook.create_sheet(_XLSX_SHEET_HELP)
    sheet.append(list(_HELP_COLUMNS))
    for column in CARD_COLUMNS:
        sheet.append(
            [
                column.header,
                column.note,
                "是" if column.required else "",
                column.example,
            ]
        )

    sheet.append([])
    sheet.append(
        [
            "环境地址怎么写",
            "一张卡片可以挂多条地址，都写在同一格「环境地址」里：每条写成 "
            "「名称 | 地址 | 模式」，多条之间用换行分隔（Excel 里按 Alt+Enter 换行），"
            "也可以直接用分号隔开。模式可写 direct / proxy（也认「直连」「代理」），"
            "省略时跟随本行的「打开方式」。\n"
            "例：生产 | https://jenkins.example.com/ ／ 测试 | http://10.0.0.9:8081/ | proxy\n"
            "这一列留空 = 这张卡片没有环境地址（导入时会清空已配置的）；"
            "如果整个文件里没有「环境地址」这一列，库里的环境地址会原样保留。",
        ]
    )
    exported_at = doc.exported_at.astimezone() if doc.exported_at else None
    sheet.append(["导出信息", f"来源 {doc.app}，配置格式 v{doc.version}，共 {doc.count} 张卡片"])
    sheet.append(
        ["导出时间", exported_at.strftime("%Y-%m-%d %H:%M:%S") if exported_at else "未记录"]
    )
    sheet.append(
        [
            "留空的含义",
            "单元格留空 = 该字段用默认值（启用=是、新标签页打开=是、校验证书=跟随全局、排序=自动）。"
            "注意这是「整行覆盖」：合并导入时，留空的列会被重置为默认值。",
        ]
    )
    sheet.append(
        [
            "怎么导入",
            "回到本工具的「管理 → 卡片管理」，点「导入配置」选这个文件。"
            "会先跑一次预演把影响列出来，确认后才真正写库。",
        ]
    )
    sheet.append(
        [
            "列名可以改吗",
            "表头也可以写成英文（slug / title / target_url / open_mode ...），"
            "多余的列会被忽略。但「标识」「名称」「目标地址」三列不能少。",
        ]
    )

    header_fill = PatternFill("solid", fgColor="E8EFFB")
    for index, width in enumerate(_HELP_WIDTHS, start=1):
        sheet.cell(row=1, column=index).fill = header_fill
        sheet.cell(row=1, column=index).font = Font(bold=True, color="1B2430")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    # 列说明之后隔一个空行才是「补充说明」区，这一区的第一列加粗当小标题
    for row_index in range(len(CARD_COLUMNS) + 3, sheet.max_row + 1):
        sheet.cell(row=row_index, column=1).font = Font(bold=True, color="1B2430")


def to_xlsx(doc: CardExportDoc) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _XLSX_SHEET_CARDS
    sheet.append([column.header for column in CARD_COLUMNS])

    for row_index, item in enumerate(doc.cards, start=2):
        for column_index, column in enumerate(CARD_COLUMNS, start=1):
            cell = sheet.cell(
                row=row_index,
                column=column_index,
                # xlsx 用换行分隔一格里的多条地址（配 wrap_text），CSV 用分号
                value=_cell_value(item, column, multi_sep="\n"),
            )
            # openpyxl 会把以「=」开头的字符串标成公式，读回来（data_only）就成了空值，
            # 名称真的写成「=汇总」这类文本时会静默丢数据，所以强制按文本写入。
            if cell.data_type == "f":
                cell.data_type = "s"

    _style_cards_sheet(sheet)
    _write_help_sheet(workbook, doc)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = [
    "CARD_COLUMNS",
    "MAX_DATA_ROWS",
    "SheetColumn",
    "SheetParseResult",
    "detect_format",
    "parse_sheet",
    "to_csv",
    "to_json",
    "to_xlsx",
]
